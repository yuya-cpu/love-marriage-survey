"""ローカル保存（レガシー）。本番保存は Google Sheets を使用。"""

from __future__ import annotations

import csv
from pathlib import Path

from response_schema import RESPONSE_HEADERS

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
RESPONSES_CSV = DATA_DIR / "responses.csv"
RESPONSES_XLSX = DATA_DIR / "responses.xlsx"


def ensure_csv_file() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not RESPONSES_CSV.exists():
        with RESPONSES_CSV.open("w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerow(RESPONSE_HEADERS)


def save_response_to_csv(record: list[str]) -> None:
    """レガシー: CSV へ1行追加"""
    ensure_csv_file()
    with RESPONSES_CSV.open("a", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerow(record)


def save_response_to_excel(record: list[str]) -> None:
    """レガシー: Excel（.xlsx）へ1行追加（openpyxl）"""
    from openpyxl import Workbook, load_workbook

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if RESPONSES_XLSX.exists():
        workbook = load_workbook(RESPONSES_XLSX)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "回答"
        sheet.append(RESPONSE_HEADERS)
    sheet.append(record)
    workbook.save(RESPONSES_XLSX)


def load_responses_from_csv() -> list[dict[str, str]]:
    ensure_csv_file()
    with RESPONSES_CSV.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))
